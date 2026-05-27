"""
ai_extractor.py
Handles document text extraction and Claude API calls for TenderEval AI.
"""

from __future__ import annotations

import io
import json
import re
from typing import Any

import anthropic


# ---------------------------------------------------------------------------
# Text extraction helpers
# ---------------------------------------------------------------------------

def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extract plain text from PDF, DOCX, XLSX or TXT bytes.
    Returns the extracted text as a single string.
    """
    name_lower = filename.lower()

    if name_lower.endswith(".pdf"):
        return _extract_pdf(file_bytes)
    elif name_lower.endswith(".docx"):
        return _extract_docx(file_bytes)
    elif name_lower.endswith((".xlsx", ".xls")):
        return _extract_xlsx(file_bytes)
    else:
        # Treat as plain text / CSV
        try:
            return file_bytes.decode("utf-8", errors="replace")
        except Exception:
            return ""


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)
    except Exception as e:
        return f"[PDF extraction error: {e}]"


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        return "\n".join(paragraphs)
    except Exception as e:
        return f"[DOCX extraction error: {e}]"


def _extract_xlsx(file_bytes: bytes) -> str:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"=== Sheet: {sheet_name} ===")
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join(str(v) if v is not None else "" for v in row)
                if row_text.strip():
                    lines.append(row_text)
        return "\n".join(lines)
    except Exception as e:
        return f"[XLSX extraction error: {e}]"


# ---------------------------------------------------------------------------
# JSON cleaning helper
# ---------------------------------------------------------------------------

def _clean_json_response(text: str) -> str:
    """Strip markdown code fences and leading/trailing whitespace."""
    text = text.strip()
    # Remove ```json ... ``` or ``` ... ```
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


# ---------------------------------------------------------------------------
# Claude API: extract and customise criteria from tender document
# ---------------------------------------------------------------------------

def extract_criteria_from_tender(
    api_key: str,
    doc_text: str,
    equipment_type: str,
    eval_name: str,
    default_criteria: dict[str, Any],
) -> dict[str, Any]:
    """
    Call Claude to review the tender document and return a customised criteria
    JSON that:
      - Keeps the default Part 1 / Part 2 / Part 3 structure
      - Adds tender-specific sub-criteria where appropriate
      - Adjusts descriptions to reference the tender requirements
      - Returns a dict matching the default_criteria structure

    Falls back to default_criteria if the API call fails.
    """
    client = anthropic.Anthropic(api_key=api_key)

    system_prompt = (
        "You are a mining procurement specialist helping to customise an equipment evaluation matrix. "
        "Your task is to review a tender document and refine the default evaluation criteria so they "
        "accurately reflect the specific requirements of this tender. "
        "Return ONLY a valid JSON object — no markdown, no commentary, no code fences."
    )

    user_prompt = f"""
You are customising an evaluation matrix for: **{eval_name}** ({equipment_type}).

## Default Criteria (JSON)
{json.dumps(default_criteria, indent=2)}

## Tender Document Text
{doc_text[:60000]}

## Instructions
1. Review the tender document and identify any specific technical requirements, mandatory features, or
   commercial conditions that should be reflected in the evaluation criteria.
2. For Part 1 (pass/fail), update "tender_reference" fields to cite specific tender clauses or page
   numbers where possible. Adjust descriptions to match tender language.
3. For Part 2 and Part 3 (SWOT), update sub-criteria descriptions to be specific to {equipment_type}.
   Add up to 3 new sub-criteria per section if the tender reveals important evaluation areas not covered.
   New sub-criteria IDs should continue the numbering (e.g. 2.11, 2.12 for section 2 additions).
4. Keep all existing IDs unchanged. Do not remove existing sub-criteria.
5. Add a "tender_reference" field to Part 1 items (string, can be empty string if not found).
6. Return the COMPLETE updated criteria JSON in exactly the same structure as the input.
   The root keys must be: "part1", "part2", "part3".
"""

    try:
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8000,
            messages=[{"role": "user", "content": user_prompt}],
            system=system_prompt,
        )
        raw = message.content[0].text
        cleaned = _clean_json_response(raw)
        result = json.loads(cleaned)
        # Ensure all three keys present
        if "part1" in result and "part2" in result and "part3" in result:
            return result
        return default_criteria
    except Exception:
        return default_criteria


# ---------------------------------------------------------------------------
# Claude API: extract SWOT scores from a bidder document
# ---------------------------------------------------------------------------

def extract_scores_from_bidder_doc(
    api_key: str,
    doc_text: str,
    bidder_name: str,
    equipment_type: str,
    eval_name: str,
    criteria: dict[str, Any],
) -> dict[str, Any]:
    """
    Call Claude to read a bidder's proposal and populate evaluation scores.

    Returns a dict:
    {
        "model_identified": str,
        "part1_scores": {
            "1.01": {"score": 1|0.5|0, "rationale": str, "tq_needed": bool}
        },
        "part2_scores": {
            "2.01": {"rating": "S"|"O"|"N"|"T"|"W", "rationale": str, "tq_needed": bool}
        },
        "part3_scores": {
            "9.01": {"rating": "S"|"O"|"N"|"T"|"W", "rationale": str, "tq_needed": bool}
        },
        "tq_list": [str],           # list of clarification questions to raise as TQs
        "key_strengths": [str],     # top 3–5 strengths
        "key_risks": [str]          # top 3–5 risks
    }
    Falls back to neutral/default scores if API call fails.
    """
    client = anthropic.Anthropic(api_key=api_key)

    system_prompt = (
        "You are a senior mining procurement engineer reviewing a vendor proposal. "
        "Your task is to objectively score a bidder against a structured evaluation matrix "
        "based solely on the content of their submitted proposal document. "
        "Be rigorous: only award Strength (S) when clear evidence exists. "
        "Flag gaps or ambiguous claims as Threats (T) or Weaknesses (W). "
        "Return ONLY a valid JSON object — no markdown, no commentary, no code fences."
    )

    # Build compact criteria summary for the prompt
    p1_summary = "\n".join(
        f"  {c['id']}: {c['name']} — {c['description']}" for c in criteria.get("part1", [])
    )
    p2_sections = []
    for sec in criteria.get("part2", []):
        sub = "\n".join(
            f"    {s['id']}: {s['name']} — {s['description']}"
            for s in sec.get("sub_criteria", [])
        )
        p2_sections.append(f"  Section {sec['section_id']} — {sec['section_name']}:\n{sub}")
    p2_summary = "\n".join(p2_sections)

    p3_sections = []
    for sec in criteria.get("part3", []):
        sub = "\n".join(
            f"    {s['id']}: {s['name']} — {s['description']}"
            for s in sec.get("sub_criteria", [])
        )
        p3_sections.append(f"  Section {sec['section_id']} — {sec['section_name']}:\n{sub}")
    p3_summary = "\n".join(p3_sections)

    user_prompt = f"""
Evaluation: **{eval_name}** — Equipment type: **{equipment_type}**
Bidder: **{bidder_name}**

## PART 1 — Compliance Gate (score each: 1=complies, 0.5=partial, 0=does not comply)
{p1_summary}

## PART 2 — Technical SWOT (rate each: S=Strength, O=Opportunity, N=Neutral, T=Threat, W=Weakness)
{p2_summary}

## PART 3 — Commercial SWOT (same S/O/N/T/W scale)
{p3_summary}

## Bidder Proposal Document
{doc_text[:70000]}

## Required JSON Output Format
{{
  "model_identified": "<equipment model name extracted from proposal, e.g. CAT 793F>",
  "part1_scores": {{
    "1.01": {{"score": 1, "rationale": "<one sentence justification>", "tq_needed": false}},
    ...all Part 1 IDs...
  }},
  "part2_scores": {{
    "2.01": {{"rating": "S", "rationale": "<one sentence justification>", "tq_needed": false}},
    ...all Part 2 sub-criteria IDs...
  }},
  "part3_scores": {{
    "9.01": {{"rating": "N", "rationale": "<one sentence justification>", "tq_needed": false}},
    ...all Part 3 sub-criteria IDs...
  }},
  "tq_list": ["<TQ question 1>", "<TQ question 2>"],
  "key_strengths": ["<strength 1>", "<strength 2>", "<strength 3>"],
  "key_risks": ["<risk 1>", "<risk 2>", "<risk 3>"]
}}

IMPORTANT: Include ALL sub-criteria IDs from the matrix. Use "N" (Neutral) and score 0.5 for any
items where the proposal is silent. Set tq_needed=true for any item where clarification is required.
"""

    # Build default fallback
    default_result = _build_default_scores(criteria)

    try:
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8000,
            messages=[{"role": "user", "content": user_prompt}],
            system=system_prompt,
        )
        raw = message.content[0].text
        cleaned = _clean_json_response(raw)
        result = json.loads(cleaned)

        # Merge with defaults to ensure all keys present
        merged = default_result.copy()
        merged.update({k: v for k, v in result.items() if v})
        return merged
    except Exception:
        return default_result


def _build_default_scores(criteria: dict[str, Any]) -> dict[str, Any]:
    """Build a neutral default score structure for all criteria."""
    p1_scores = {
        c["id"]: {"score": 0.5, "rationale": "Not yet reviewed.", "tq_needed": False}
        for c in criteria.get("part1", [])
    }
    p2_scores = {}
    for sec in criteria.get("part2", []):
        for sub in sec.get("sub_criteria", []):
            p2_scores[sub["id"]] = {"rating": "N", "rationale": "Not yet reviewed.", "tq_needed": False}
    p3_scores = {}
    for sec in criteria.get("part3", []):
        for sub in sec.get("sub_criteria", []):
            p3_scores[sub["id"]] = {"rating": "N", "rationale": "Not yet reviewed.", "tq_needed": False}

    return {
        "model_identified": "Unknown",
        "part1_scores": p1_scores,
        "part2_scores": p2_scores,
        "part3_scores": p3_scores,
        "tq_list": [],
        "key_strengths": [],
        "key_risks": [],
    }
