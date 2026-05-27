"""
eval_sheet.py
Export a fillable evaluator workbook and parse completed sheets back into scores.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

_NAVY = "052B48"
_ORANGE = "D05F27"
_FILL_P1 = "DCE6F1"   # light blue  – Part 1
_FILL_P2 = "E2EFDA"   # light green – Part 2
_FILL_P3 = "FFF2CC"   # light yellow – Part 3
_FILL_EDIT = "FFFFFF"  # white – editable cells
_THIN = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_EDITABLE_COLS = {5, 6, 7, 8}  # Score/Rating, Rationale, Evaluator Name, Notes


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def generate_evaluator_sheet(eval_data: dict[str, Any]) -> bytes:
    """
    Return bytes of an .xlsx workbook with one sheet per bidder.
    Each sheet lists every criterion with editable Score/Rating, Rationale,
    Evaluator Name, and Notes columns.  Existing AI scores are pre-filled.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    bidders = eval_data.get("bidders", [])
    criteria = eval_data.get("criteria", {})
    scores = eval_data.get("scores", {})
    eval_name = eval_data.get("eval_name", "TenderEval")
    equipment_type = eval_data.get("equipment_type", "")

    for bidder in bidders:
        ws = wb.create_sheet(title=bidder[:31])
        _write_bidder_sheet(
            ws, bidder, criteria, scores.get(bidder, {}), eval_name, equipment_type
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_bidder_sheet(
    ws,
    bidder: str,
    criteria: dict,
    b_scores: dict,
    eval_name: str,
    equipment_type: str,
) -> None:
    # Column widths
    for col, width in zip("ABCDEFGH", [12, 8, 28, 45, 14, 45, 20, 35]):
        ws.column_dimensions[col].width = width

    # Row 1 – title
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"{eval_name}  |  {equipment_type}  |  Evaluator Sheet — {bidder}"
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=_NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Row 2 – instructions
    ws.merge_cells("A2:H2")
    c = ws["A2"]
    c.value = (
        "Part 1 Score: 1 = Complies  |  0.5 = Partial  |  0 = Does not comply    "
        "Parts 2 & 3 Rating: S = Strength  |  O = Opportunity  |  N = Neutral  |  "
        "T = Threat  |  W = Weakness    Fill in Score/Rating, Rationale, and your name, then return this file."
    )
    c.font = Font(italic=True, size=9, color="595959")
    c.fill = PatternFill("solid", fgColor="FFF2CC")
    c.alignment = Alignment(wrap_text=True, horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 30

    # Row 3 – headers
    headers = ["Part", "ID", "Criterion", "Description", "Score / Rating",
               "Rationale", "Evaluator Name", "Notes"]
    for col_idx, label in enumerate(headers, 1):
        c = ws.cell(row=3, column=col_idx, value=label)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=_ORANGE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 18
    ws.freeze_panes = "A4"

    # Data-validation dropdowns
    dv_p1 = DataValidation(
        type="list", formula1='"1,0.5,0"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid score",
        error="Enter 1, 0.5, or 0",
    )
    dv_p23 = DataValidation(
        type="list", formula1='"S,O,N,T,W"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid rating",
        error="Enter S, O, N, T, or W",
    )
    ws.add_data_validation(dv_p1)
    ws.add_data_validation(dv_p23)

    row = 4

    def write_row(part_label, crit_id, name, desc, score_val, rationale, fill_color, dv):
        nonlocal row
        data = [part_label, crit_id, name, desc, score_val, rationale, "", ""]
        for col_idx, val in enumerate(data, 1):
            c = ws.cell(row=row, column=col_idx, value=val)
            c.border = _BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = PatternFill("solid", fgColor=_FILL_EDIT if col_idx in _EDITABLE_COLS else fill_color)
        ws.row_dimensions[row].height = 40
        dv.add(ws.cell(row=row, column=5))
        row += 1

    # Part 1
    p1 = b_scores.get("part1_scores", {})
    for c in criteria.get("part1", []):
        entry = p1.get(c["id"], {})
        raw = entry.get("score", "")
        try:
            score_val = float(raw) if raw != "" else ""
        except (TypeError, ValueError):
            score_val = ""
        write_row(
            "Part 1 – Compliance", c["id"], c["name"], c.get("description", ""),
            score_val, entry.get("rationale", ""), _FILL_P1, dv_p1,
        )

    # Part 2
    p2 = b_scores.get("part2_scores", {})
    for sec in criteria.get("part2", []):
        for sub in sec.get("sub_criteria", []):
            entry = p2.get(sub["id"], {})
            write_row(
                f"Part 2 – {sec['section_name']}", sub["id"], sub["name"],
                sub.get("description", ""), entry.get("rating", ""),
                entry.get("rationale", ""), _FILL_P2, dv_p23,
            )

    # Part 3
    p3 = b_scores.get("part3_scores", {})
    for sec in criteria.get("part3", []):
        for sub in sec.get("sub_criteria", []):
            entry = p3.get(sub["id"], {})
            write_row(
                f"Part 3 – {sec['section_name']}", sub["id"], sub["name"],
                sub.get("description", ""), entry.get("rating", ""),
                entry.get("rationale", ""), _FILL_P3, dv_p23,
            )


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def parse_evaluator_sheet(
    file_bytes: bytes,
    bidder_name: str,
    criteria: dict[str, Any],
) -> dict[str, Any]:
    """
    Read a completed evaluator workbook and return a partial scores dict
    for one bidder:  {part1_scores, part2_scores, part3_scores}.

    Finds the sheet whose name best matches bidder_name.
    Returns {} if no matching sheet is found.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    ws = _find_sheet(wb, bidder_name)
    if ws is None:
        return {}

    # Build ID lookup sets
    p1_ids = {c["id"] for c in criteria.get("part1", [])}
    p2_ids = {
        sub["id"]
        for sec in criteria.get("part2", [])
        for sub in sec.get("sub_criteria", [])
    }
    p3_ids = {
        sub["id"]
        for sec in criteria.get("part3", [])
        for sub in sec.get("sub_criteria", [])
    }

    p1_scores: dict = {}
    p2_scores: dict = {}
    p3_scores: dict = {}

    for row in ws.iter_rows(min_row=4, values_only=True):
        crit_id = str(row[1]).strip() if row[1] is not None else ""
        if not crit_id or crit_id.lower() == "none":
            continue

        score_raw = row[4]
        rationale = str(row[5]).strip() if row[5] is not None else ""
        evaluator = str(row[6]).strip() if row[6] is not None else ""
        notes = str(row[7]).strip() if row[7] is not None else ""

        entry_base = {
            "rationale": rationale,
            "tq_needed": False,
            "changed_by": evaluator,
            "notes": notes,
        }

        if crit_id in p1_ids:
            try:
                score = float(score_raw) if score_raw not in (None, "") else None
                if score is not None:
                    score = max(0.0, min(1.0, score))
            except (ValueError, TypeError):
                score = None
            if score is not None:
                p1_scores[crit_id] = {**entry_base, "score": score}

        elif crit_id in p2_ids:
            rating = str(score_raw).strip().upper() if score_raw not in (None, "") else ""
            if rating in ("S", "O", "N", "T", "W"):
                p2_scores[crit_id] = {**entry_base, "rating": rating}

        elif crit_id in p3_ids:
            rating = str(score_raw).strip().upper() if score_raw not in (None, "") else ""
            if rating in ("S", "O", "N", "T", "W"):
                p3_scores[crit_id] = {**entry_base, "rating": rating}

    return {
        "part1_scores": p1_scores,
        "part2_scores": p2_scores,
        "part3_scores": p3_scores,
    }


def list_sheet_bidders(file_bytes: bytes) -> list[str]:
    """Return the sheet names found in an uploaded evaluator workbook."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _find_sheet(wb, bidder_name: str):
    """Return the worksheet whose name best matches bidder_name, or None."""
    target = bidder_name.strip().lower()
    for name in wb.sheetnames:
        if name.strip().lower() == target[:31]:
            return wb[name]
    for name in wb.sheetnames:
        if target in name.strip().lower() or name.strip().lower() in target:
            return wb[name]
    return None
