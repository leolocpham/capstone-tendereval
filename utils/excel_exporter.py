"""
excel_exporter.py
Generates the evaluation Excel workbook for TenderEval AI.
Sheets: Comparison Summary | Assessment (per bidder) | TCO Analysis
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Brand colours
# ---------------------------------------------------------------------------
NAVY = "052B48"
ORANGE = "D05F27"
WHITE = "FFFFFF"
LIGHT_GREY = "F2F2F2"
MID_GREY = "BFBFBF"
DARK_GREY = "595959"

# SWOT colours
SWOT_FILLS = {
    "S": PatternFill("solid", fgColor="00B050"),   # green
    "O": PatternFill("solid", fgColor="FFFF00"),   # yellow
    "N": PatternFill("solid", fgColor="BFBFBF"),   # grey
    "T": PatternFill("solid", fgColor="FFC000"),   # light orange
    "W": PatternFill("solid", fgColor="FF0000"),   # red
}
SWOT_FONT_DARK = {"S", "N", "T"}  # use dark font for readability

PASS_FILL = PatternFill("solid", fgColor="00B050")
FAIL_FILL = PatternFill("solid", fgColor="FF0000")
PARTIAL_FILL = PatternFill("solid", fgColor="FFC000")

NAVY_FILL = PatternFill("solid", fgColor=NAVY)
ORANGE_FILL = PatternFill("solid", fgColor=ORANGE)
GREY_FILL = PatternFill("solid", fgColor=LIGHT_GREY)


# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------

def _hdr_font(bold=True, size=11, color=WHITE):
    return Font(name="Arial", bold=bold, size=size, color=color)


def _body_font(bold=False, size=10, color="000000"):
    return Font(name="Arial", bold=bold, size=size, color=color)


def _thin_border():
    side = Side(style="thin")
    return Border(left=side, right=side, top=side, bottom=side)


def _set_col_width(ws, col_letter: str, width: float):
    ws.column_dimensions[col_letter].width = width


def _write_header_row(ws, row: int, values: list, fill, font=None, height=20):
    if font is None:
        font = _hdr_font()
    ws.row_dimensions[row].height = height
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()


def _write_cell(ws, row, col, value, fill=None, font=None, align="left", wrap=False, border=True):
    cell = ws.cell(row=row, column=col, value=value)
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    else:
        cell.font = _body_font()
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border:
        cell.border = _thin_border()
    return cell


# ---------------------------------------------------------------------------
# Main export entry point
# ---------------------------------------------------------------------------

def generate_assessment_matrix(eval_data: dict[str, Any]) -> bytes:
    """
    Build the full evaluation Excel workbook and return as bytes.

    eval_data keys used:
        eval_name, equipment_type, tender_number,
        bidders, criteria, scores, tco, tco_results
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    bidders: list[str] = eval_data.get("bidders", [])
    criteria: dict = eval_data.get("criteria", {})
    scores: dict = eval_data.get("scores", {})
    tco_results: dict = eval_data.get("tco_results", {})

    # --- Build sheets ---
    _build_comparison_summary(wb, eval_data, bidders, criteria, scores, tco_results)
    for bidder in bidders:
        _build_bidder_assessment(wb, eval_data, bidder, criteria, scores.get(bidder, {}))
    _build_tco_sheet(wb, eval_data, bidders, tco_results)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# Sheet 1: Comparison Summary
# ---------------------------------------------------------------------------

def _build_comparison_summary(wb, eval_data, bidders, criteria, scores, tco_results):
    ws = wb.create_sheet("Comparison Summary")
    ws.sheet_view.showGridLines = False

    eval_name = eval_data.get("eval_name", "Evaluation")
    equipment_type = eval_data.get("equipment_type", "")
    tender_number = eval_data.get("tender_number", "")

    row = 1
    # Title block
    ws.merge_cells(f"A{row}:G{row}")
    title_cell = ws.cell(row=row, column=1, value=f"TenderEval AI — {eval_name}")
    title_cell.fill = NAVY_FILL
    title_cell.font = Font(name="Arial", bold=True, size=14, color=WHITE)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 28
    row += 1

    ws.merge_cells(f"A{row}:G{row}")
    sub_cell = ws.cell(row=row, column=1,
                       value=f"Equipment: {equipment_type}  |  Tender: {tender_number}")
    sub_cell.fill = ORANGE_FILL
    sub_cell.font = _hdr_font(size=10)
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18
    row += 2

    # --- PART 1 SUMMARY ---
    ws.merge_cells(f"A{row}:G{row}")
    ws.cell(row=row, column=1, value="PART 1 — PROPOSAL COMPLIANCE (PASS / FAIL GATE)").fill = NAVY_FILL
    ws.cell(row=row, column=1).font = _hdr_font()
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18
    row += 1

    headers = ["#", "Criterion", "Weight"] + bidders
    _write_header_row(ws, row, headers, NAVY_FILL)
    row += 1

    from utils.tco import calculate_part1_score, part1_pass

    p1_criteria = criteria.get("part1", [])
    for c in p1_criteria:
        ws.row_dimensions[row].height = 28
        _write_cell(ws, row, 1, c["id"], align="center")
        _write_cell(ws, row, 2, c["name"], wrap=True)
        _write_cell(ws, row, 3, f"{c.get('weight', 0):.0%}", align="center")
        for b_idx, bidder in enumerate(bidders, 4):
            b_scores = scores.get(bidder, {})
            entry = b_scores.get("part1_scores", {}).get(c["id"], {})
            score_val = entry.get("score", 0.5)
            label = {1: "PASS", 0.5: "PARTIAL", 0: "FAIL"}.get(score_val, str(score_val))
            fill = {1: PASS_FILL, 0.5: PARTIAL_FILL, 0: FAIL_FILL}.get(score_val, GREY_FILL)
            cell = _write_cell(ws, row, b_idx, label, fill=fill, align="center")
            cell.font = Font(name="Arial", bold=True, size=10, color=WHITE if score_val == 0 else "000000")
        row += 1

    # Part 1 totals
    ws.row_dimensions[row].height = 20
    _write_cell(ws, row, 1, "", fill=GREY_FILL)
    _write_cell(ws, row, 2, "WEIGHTED COMPLIANCE SCORE", fill=GREY_FILL,
                font=_body_font(bold=True))
    _write_cell(ws, row, 3, "", fill=GREY_FILL)
    for b_idx, bidder in enumerate(bidders, 4):
        b_scores = scores.get(bidder, {})
        overall = calculate_part1_score(b_scores.get("part1_scores", {}), p1_criteria)
        passed = part1_pass(overall)
        fill = PASS_FILL if passed else FAIL_FILL
        label = f"{overall:.1%}  {'✓ PASS' if passed else '✗ FAIL'}"
        cell = _write_cell(ws, row, b_idx, label, fill=fill, align="center",
                           font=_body_font(bold=True, color=WHITE if not passed else "000000"))
    row += 2

    # --- PART 2 + 3 SWOT SUMMARY ---
    from utils.tco import calculate_section_score, calculate_part_score

    for part_key, part_label, score_key in [
        ("part2", "PART 2 — TECHNICAL SWOT", "part2_scores"),
        ("part3", "PART 3 — COMMERCIAL SWOT", "part3_scores"),
    ]:
        ws.merge_cells(f"A{row}:G{row}")
        hdr_cell = ws.cell(row=row, column=1, value=part_label)
        hdr_cell.fill = NAVY_FILL
        hdr_cell.font = _hdr_font()
        hdr_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 18
        row += 1

        sections = criteria.get(part_key, [])
        _write_header_row(ws, row, ["§", "Section", "Weight"] + bidders, NAVY_FILL)
        row += 1

        for sec in sections:
            ws.row_dimensions[row].height = 20
            _write_cell(ws, row, 1, f"§{sec['section_id']}", align="center")
            _write_cell(ws, row, 2, sec["section_name"])
            _write_cell(ws, row, 3, f"{sec.get('weight', 1.0):.1f}×", align="center")
            for b_idx, bidder in enumerate(bidders, 4):
                b_scores = scores.get(bidder, {}).get(score_key, {})
                raw = calculate_section_score(sec["section_id"], b_scores, sections)
                weighted = raw * sec.get("weight", 1.0)
                fill = PASS_FILL if weighted > 0 else (FAIL_FILL if weighted < 0 else GREY_FILL)
                _write_cell(ws, row, b_idx, f"{weighted:+.1f}", fill=fill, align="center",
                            font=_body_font(bold=True))
            row += 1

        # Part total
        ws.row_dimensions[row].height = 20
        _write_cell(ws, row, 2, f"{part_label.split('—')[1].strip()} TOTAL",
                    fill=GREY_FILL, font=_body_font(bold=True))
        _write_cell(ws, row, 1, "", fill=GREY_FILL)
        _write_cell(ws, row, 3, "", fill=GREY_FILL)
        for b_idx, bidder in enumerate(bidders, 4):
            b_scores = scores.get(bidder, {}).get(score_key, {})
            total = calculate_part_score(b_scores, sections)
            fill = PASS_FILL if total > 0 else (FAIL_FILL if total < 0 else GREY_FILL)
            _write_cell(ws, row, b_idx, f"{total:+.2f}", fill=fill, align="center",
                        font=_body_font(bold=True))
        row += 2

    # --- TCO SUMMARY ---
    ws.merge_cells(f"A{row}:G{row}")
    ws.cell(row=row, column=1, value="TCO SUMMARY").fill = NAVY_FILL
    ws.cell(row=row, column=1).font = _hdr_font()
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18
    row += 1

    tco_rows = [
        ("Total CAPEX (AUD)", "total_capex"),
        ("Total OPEX (AUD)", "total_opex"),
        ("Total TCO (AUD)", "total_tco"),
        ("NPV @ 8% (AUD)", "npv"),
        ("Cost per Operating Hour (AUD/hr)", "cost_per_hour"),
    ]
    _write_header_row(ws, row, ["", "Metric", ""] + bidders, NAVY_FILL)
    row += 1
    for label, key in tco_rows:
        ws.row_dimensions[row].height = 18
        _write_cell(ws, row, 1, "")
        _write_cell(ws, row, 2, label)
        _write_cell(ws, row, 3, "")
        for b_idx, bidder in enumerate(bidders, 4):
            val = tco_results.get(bidder, {}).get(key, 0)
            fmt = f"${val:,.0f}" if "Hour" not in label else f"${val:,.2f}"
            _write_cell(ws, row, b_idx, fmt, align="right")
        row += 1

    # Column widths
    _set_col_width(ws, "A", 8)
    _set_col_width(ws, "B", 38)
    _set_col_width(ws, "C", 10)
    for i in range(len(bidders)):
        _set_col_width(ws, get_column_letter(4 + i), 22)


# ---------------------------------------------------------------------------
# Sheet per bidder: Assessment
# ---------------------------------------------------------------------------

def _build_bidder_assessment(wb, eval_data, bidder, criteria, bidder_scores):
    sheet_name = f"Assessment – {bidder[:25]}"
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    eval_name = eval_data.get("eval_name", "")
    equipment_type = eval_data.get("equipment_type", "")
    model_id = bidder_scores.get("model_identified", "")

    row = 1
    # Header block
    ws.merge_cells(f"A{row}:F{row}")
    ws.cell(row=row, column=1,
            value=f"Assessment: {bidder}  |  {eval_name}  |  {equipment_type}  |  Model: {model_id}").fill = NAVY_FILL
    ws.cell(row=row, column=1).font = _hdr_font(size=12)
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 26
    row += 2

    # ---- PART 1 ----
    ws.merge_cells(f"A{row}:F{row}")
    ws.cell(row=row, column=1, value="PART 1 — PROPOSAL COMPLIANCE").fill = NAVY_FILL
    ws.cell(row=row, column=1).font = _hdr_font()
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18
    row += 1

    _write_header_row(ws, row, ["#", "Criterion", "Description", "Weight", "Score", "AI Rationale"],
                      NAVY_FILL)
    row += 1

    p1_scores = bidder_scores.get("part1_scores", {})
    for c in criteria.get("part1", []):
        ws.row_dimensions[row].height = 38
        entry = p1_scores.get(c["id"], {})
        score_val = entry.get("score", 0.5)
        label = {1: "PASS", 0.5: "PARTIAL", 0: "FAIL"}.get(score_val, str(score_val))
        fill = {1: PASS_FILL, 0.5: PARTIAL_FILL, 0: FAIL_FILL}.get(score_val, GREY_FILL)
        _write_cell(ws, row, 1, c["id"], align="center")
        _write_cell(ws, row, 2, c["name"])
        _write_cell(ws, row, 3, c.get("description", ""), wrap=True)
        _write_cell(ws, row, 4, f"{c.get('weight', 0):.0%}", align="center")
        _write_cell(ws, row, 5, label, fill=fill, align="center",
                    font=_body_font(bold=True, color=WHITE if score_val == 0 else "000000"))
        _write_cell(ws, row, 6, entry.get("rationale", ""), wrap=True)
        row += 1

    row += 1

    # ---- PARTS 2 & 3 ----
    for part_key, part_title, score_key in [
        ("part2", "PART 2 — QUALITATIVE TECHNICAL (SWOT)", "part2_scores"),
        ("part3", "PART 3 — QUALITATIVE COMMERCIAL (SWOT)", "part3_scores"),
    ]:
        ws.merge_cells(f"A{row}:F{row}")
        ws.cell(row=row, column=1, value=part_title).fill = NAVY_FILL
        ws.cell(row=row, column=1).font = _hdr_font()
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 18
        row += 1

        part_scores = bidder_scores.get(score_key, {})
        for sec in criteria.get(part_key, []):
            # Section header
            ws.merge_cells(f"A{row}:F{row}")
            sec_cell = ws.cell(row=row, column=1,
                               value=f"§{sec['section_id']}  {sec['section_name']}  (Weight: {sec.get('weight', 1.0):.1f}×)")
            sec_cell.fill = ORANGE_FILL
            sec_cell.font = _hdr_font(size=10)
            sec_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row].height = 18
            row += 1

            _write_header_row(ws, row, ["#", "Sub-Criterion", "Description", "", "Rating", "AI Rationale"],
                              PatternFill("solid", fgColor=DARK_GREY))
            row += 1

            for sub in sec.get("sub_criteria", []):
                ws.row_dimensions[row].height = 38
                entry = part_scores.get(sub["id"], {})
                rating = entry.get("rating", "N").upper()
                fill = SWOT_FILLS.get(rating, GREY_FILL)
                font_color = "000000" if rating in SWOT_FONT_DARK else "000000"
                _write_cell(ws, row, 1, sub["id"], align="center")
                _write_cell(ws, row, 2, sub["name"])
                _write_cell(ws, row, 3, sub.get("description", ""), wrap=True)
                _write_cell(ws, row, 4, "")
                _write_cell(ws, row, 5, rating, fill=fill, align="center",
                            font=_body_font(bold=True, color=font_color))
                _write_cell(ws, row, 6, entry.get("rationale", ""), wrap=True)
                row += 1
            row += 1

    # Key strengths / risks
    _write_cell(ws, row, 1, "KEY STRENGTHS", fill=NAVY_FILL,
                font=_hdr_font(), align="left")
    ws.merge_cells(f"B{row}:F{row}")
    ws.row_dimensions[row].height = 18
    row += 1
    for s in bidder_scores.get("key_strengths", []):
        ws.merge_cells(f"A{row}:F{row}")
        _write_cell(ws, row, 1, f"✓  {s}", fill=GREY_FILL, wrap=True)
        ws.row_dimensions[row].height = 22
        row += 1
    row += 1

    _write_cell(ws, row, 1, "KEY RISKS", fill=NAVY_FILL,
                font=_hdr_font(), align="left")
    ws.merge_cells(f"B{row}:F{row}")
    ws.row_dimensions[row].height = 18
    row += 1
    for r_item in bidder_scores.get("key_risks", []):
        ws.merge_cells(f"A{row}:F{row}")
        _write_cell(ws, row, 1, f"⚠  {r_item}", fill=GREY_FILL, wrap=True)
        ws.row_dimensions[row].height = 22
        row += 1

    # Column widths
    _set_col_width(ws, "A", 9)
    _set_col_width(ws, "B", 28)
    _set_col_width(ws, "C", 40)
    _set_col_width(ws, "D", 4)
    _set_col_width(ws, "E", 10)
    _set_col_width(ws, "F", 50)


# ---------------------------------------------------------------------------
# Sheet: TCO Analysis
# ---------------------------------------------------------------------------

def _build_tco_sheet(wb, eval_data, bidders, tco_results):
    ws = wb.create_sheet("TCO Analysis")
    ws.sheet_view.showGridLines = False

    tco_inputs_all: dict = eval_data.get("tco", {})

    row = 1
    ws.merge_cells(f"A{row}:{get_column_letter(2 + len(bidders))}{row}")
    ws.cell(row=row, column=1, value="TCO ANALYSIS").fill = NAVY_FILL
    ws.cell(row=row, column=1).font = _hdr_font(size=13)
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 26
    row += 2

    # Header row
    _write_header_row(ws, row, ["#", "Item"] + bidders, NAVY_FILL)
    row += 1

    capex_rows = [
        ("C1", "Ready-to-Work (RTW) Cost", "rtw_cost"),
        ("C2", "Change Management CAPEX", "change_mgmt_capex"),
        ("C3", "Training Cost", "training"),
        ("C4", "Seed Components", "seed_components"),
    ]

    ws.merge_cells(f"A{row}:{get_column_letter(2 + len(bidders))}{row}")
    ws.cell(row=row, column=1, value="CAPITAL EXPENDITURE (CAPEX)").fill = ORANGE_FILL
    ws.cell(row=row, column=1).font = _hdr_font(size=10)
    ws.row_dimensions[row].height = 16
    row += 1

    for item_id, label, key in capex_rows:
        ws.row_dimensions[row].height = 18
        _write_cell(ws, row, 1, item_id, align="center")
        _write_cell(ws, row, 2, label)
        for b_idx, bidder in enumerate(bidders, 3):
            val = tco_inputs_all.get(bidder, {}).get(key, 0)
            _write_cell(ws, row, b_idx, val, align="right",
                        font=_body_font())
            ws.cell(row=row, column=b_idx).number_format = "$#,##0"
        row += 1

    # Total CAPEX
    ws.row_dimensions[row].height = 18
    _write_cell(ws, row, 1, "", fill=GREY_FILL)
    _write_cell(ws, row, 2, "TOTAL CAPEX", fill=GREY_FILL, font=_body_font(bold=True))
    for b_idx, bidder in enumerate(bidders, 3):
        val = tco_results.get(bidder, {}).get("total_capex", 0)
        c = _write_cell(ws, row, b_idx, val, align="right", fill=GREY_FILL,
                        font=_body_font(bold=True))
        c.number_format = "$#,##0"
    row += 2

    # OPEX per year
    ws.merge_cells(f"A{row}:{get_column_letter(2 + len(bidders))}{row}")
    ws.cell(row=row, column=1, value="OPERATING EXPENDITURE (OPEX) BY YEAR").fill = ORANGE_FILL
    ws.cell(row=row, column=1).font = _hdr_font(size=10)
    ws.row_dimensions[row].height = 16
    row += 1

    _write_header_row(ws, row, ["Year", "Component"] + bidders,
                      PatternFill("solid", fgColor=DARK_GREY))
    row += 1

    max_years = max(
        (len(tco_inputs_all.get(b, {}).get("opex_years", [])) for b in bidders),
        default=1,
    )

    for yr_idx in range(max_years):
        for comp_label, comp_key in [
            ("Operating", "operating"),
            ("Maintenance", "maintenance"),
            ("Indirect", "indirect"),
        ]:
            ws.row_dimensions[row].height = 16
            _write_cell(ws, row, 1, f"Year {yr_idx + 1}", align="center")
            _write_cell(ws, row, 2, comp_label)
            for b_idx, bidder in enumerate(bidders, 3):
                opex_years = tco_inputs_all.get(bidder, {}).get("opex_years", [])
                val = opex_years[yr_idx].get(comp_key, 0) if yr_idx < len(opex_years) else 0
                c = _write_cell(ws, row, b_idx, val, align="right")
                c.number_format = "$#,##0"
            row += 1
    row += 1

    # Summary block
    ws.merge_cells(f"A{row}:{get_column_letter(2 + len(bidders))}{row}")
    ws.cell(row=row, column=1, value="TCO SUMMARY").fill = NAVY_FILL
    ws.cell(row=row, column=1).font = _hdr_font()
    ws.row_dimensions[row].height = 18
    row += 1

    for label, key, fmt in [
        ("Total CAPEX (AUD)", "total_capex", "$#,##0"),
        ("Total OPEX (AUD)", "total_opex", "$#,##0"),
        ("Total TCO (AUD)", "total_tco", "$#,##0"),
        ("NPV @ 8% (AUD)", "npv", "$#,##0"),
        ("Cost per Operating Hour (AUD/hr)", "cost_per_hour", "$#,##0.00"),
    ]:
        ws.row_dimensions[row].height = 18
        _write_cell(ws, row, 1, "", fill=GREY_FILL)
        _write_cell(ws, row, 2, label, fill=GREY_FILL, font=_body_font(bold=True))
        for b_idx, bidder in enumerate(bidders, 3):
            val = tco_results.get(bidder, {}).get(key, 0)
            c = _write_cell(ws, row, b_idx, val, fill=GREY_FILL, align="right",
                            font=_body_font(bold=True))
            c.number_format = fmt
        row += 1

    # Column widths
    _set_col_width(ws, "A", 9)
    _set_col_width(ws, "B", 38)
    for i in range(len(bidders)):
        _set_col_width(ws, get_column_letter(3 + i), 22)
